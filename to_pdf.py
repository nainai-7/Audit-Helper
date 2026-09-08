# -*- coding: utf-8 -*-
"""Excel 按列拆分 PDF 工具（通用版，带图形界面，可打包成 exe 分享）

功能：读取 Excel 某个工作表，按指定列的值把行分组，每组生成一个 PDF 表格文件。
      表头从表头行自动读取，支持指定合计列（可多列）、横/纵向纸张、自定义标题。

"""
import os
import re
import ctypes
import threading
import queue
from collections import OrderedDict
from datetime import datetime, date

import openpyxl
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import getAscent, getDescent
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas as pdfcanvas

# ---------------------------------------------------------------- 字体注册
# 优先用系统宋体/黑体；找不到时退回 reportlab 内置 CID 字体（保证 exe 在别的机器能跑）
FONT = "SimSun"
FONT_BOLD = "HeiTi"


def _register_fonts():
    global FONT, FONT_BOLD
    try:
        pdfmetrics.registerFont(TTFont("SimSun", r"C:\Windows\Fonts\simsun.ttc", subfontIndex=0))
    except Exception:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        FONT = "STSong-Light"
    try:
        pdfmetrics.registerFont(TTFont("HeiTi", r"C:\Windows\Fonts\simhei.ttf"))
    except Exception:
        FONT_BOLD = FONT  # 找不到黑体时表头不加粗


_register_fonts()

# ---------------------------------------------------------------- 通用小函数


def col_to_index(s):
    """'B' / 'b' / '2' -> 1 开始的列号"""
    s = str(s).strip().upper()
    if not s:
        raise ValueError("列号不能为空")
    if s.isdigit():
        n = int(s)
        if n < 1:
            raise ValueError("列号从 1 开始")
        return n
    n = 0
    for ch in s:
        if not ("A" <= ch <= "Z"):
            raise ValueError(f"无法识别的列号: {s!r}")
        n = n * 26 + ord(ch) - 64
    return n


def index_to_col(n):
    """1 -> 'A', 27 -> 'AA'"""
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def sanitize_filename(name):
    name = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", str(name)).strip().strip(".")
    return name or "未命名"


def fmt_cell(v):
    """单元格显示文本：空->空串；整数浮点->整数；日期->2026-08-31；其余->str"""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def fmt_num(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def wrap(text, size, maxw, font=None):
    if not text:
        return [""]
    font = font or FONT
    lines, cur = [], ""
    for ch in text:
        if pdfmetrics.stringWidth(cur + ch, font, size) <= maxw:
            cur += ch
        else:
            lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines or [""]


# ---------------------------------------------------------------- 核心拆分逻辑

PAD = 2.5
MIN_COL_W = 40.0
TITLE_H = 28
TOP_MARGIN = 46
BOTTOM_MARGIN = 28
SIDE_MARGIN = 30


def build_pdf(file_name, headers, rows, sum_rels, out_dir, orientation,
              title_text, sum_label, size):
    """生成一个 PDF 表格，返回 (路径, 页数)。
    headers: 表头文字列表；rows: 原始值行列表；sum_rels: 合计列的相对下标(0 开始)。
    """
    page_w, page_h = landscape(A4) if orientation == "landscape" else A4
    table_w = page_w - SIDE_MARGIN * 2
    x0 = SIDE_MARGIN
    ncols = len(headers)

    # 合计值
    totals = [0.0] * ncols
    has_num = [False] * ncols
    for r in rows:
        for i in sum_rels:
            v = r[i] if i < len(r) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                totals[i] += v
                has_num[i] = True

    # ---- 自适应列宽
    needs = []
    for i, h in enumerate(headers):
        wmax = pdfmetrics.stringWidth(h, FONT_BOLD, size)
        for r in rows:
            v = fmt_cell(r[i] if i < len(r) else None)
            w = pdfmetrics.stringWidth(v, FONT, size)
            if w > wmax:
                wmax = w
        needs.append(max(wmax + PAD * 2, MIN_COL_W))
    total_need = sum(needs)
    if total_need <= table_w:
        widths = needs[:]
        j = max(range(ncols), key=lambda k: widths[k])
        widths[j] += table_w - total_need  # 多余空间给最宽列
    else:
        widths = [max(MIN_COL_W, w * table_w / total_need) for w in needs]
        diff = table_w - sum(widths)
        j = max(range(ncols), key=lambda k: widths[k])
        widths[j] += diff

    lh = size * 1.25
    header_h = lh + PAD * 2

    def cell_lines(values):
        return [wrap(fmt_cell(v), size, w - 4) for v, w in zip(values, widths)]

    # ---- 行高
    heights = []
    for r in rows:
        n = max(len(x) for x in cell_lines(r))
        heights.append(max(n * lh + PAD * 2, size + PAD * 2 + 2))

    path = os.path.join(out_dir, sanitize_filename(file_name) + ".pdf")
    c = pdfcanvas.Canvas(path, pagesize=(page_w, page_h))
    page_count = 1
    c.setTitle(file_name)

    def vcenter_baseline(y, height, n):
        ascent = getAscent(FONT, size)
        descent = -getDescent(FONT, size)
        text_h = (n - 1) * lh + ascent + descent
        return y + (height + text_h) / 2 - ascent

    def draw_title():
        if title_text:
            c.setFont(FONT_BOLD, size)
            c.drawString(x0, page_h - TITLE_H, title_text)

    def draw_head(y):
        c.setFont(FONT_BOLD, size)
        c.rect(x0, y, table_w, header_h, stroke=1, fill=0)
        x = x0
        for i, (h, w) in enumerate(zip(headers, widths)):
            if i < ncols - 1:
                c.line(x + w, y, x + w, y + header_h)
            c.drawCentredString(x + w / 2, vcenter_baseline(y, header_h, 1), h)
            x += w
        return y - header_h

    def draw_row(y, values, height):
        lines = cell_lines(values)
        c.setFont(FONT, size)
        c.rect(x0, y, table_w, height, stroke=1, fill=0)
        x = x0
        for i, w in enumerate(widths):
            if i < ncols - 1:
                c.line(x + w, y, x + w, y + height)
            cell = lines[i]
            n = len(cell)
            raw = values[i] if i < len(values) else None
            # 数值或合计列右对齐，其余左对齐
            right = (i in sum_rels) or isinstance(raw, (int, float)) and not isinstance(raw, bool)
            ty = vcenter_baseline(y, height, n)
            for ln in cell:
                if right:
                    c.drawRightString(x + w - 2, ty, ln)
                else:
                    c.drawString(x + 2, ty, ln)
                ty -= lh
            x += w
        return y - height

    def draw_page_header():
        draw_title()
        return draw_head(page_h - TOP_MARGIN)

    # ---- 数据行
    y = draw_page_header()
    for r, rh in zip(rows, heights):
        if y - rh < BOTTOM_MARGIN:
            c.showPage()
            page_count += 1
            y = draw_page_header()
        y = draw_row(y, r, rh)

    # ---- 汇总行（有合计列才画）
    if sum_rels:
        if y - header_h < BOTTOM_MARGIN:
            c.showPage()
            page_count += 1
            y = draw_page_header()
        c.setFont(FONT, size)
        c.rect(x0, y, table_w, header_h, stroke=1, fill=0)
        x = x0
        for i, w in enumerate(widths):
            if i < ncols - 1:
                c.line(x + w, y, x + w, y + header_h)
            if i == 0:
                c.drawString(x + 2, vcenter_baseline(y, header_h, 1), sum_label)
            elif i in sum_rels and has_num[i]:
                c.drawRightString(x + w - 2, vcenter_baseline(y, header_h, 1), fmt_num(totals[i]))
            x += w

    c.save()
    return path, page_count


def run_split(cfg, log=print):
    """执行拆分。cfg 各键见 GUI 收集逻辑。返回 (成功数, 失败列表, 多页数, 文件数)。"""
    src = cfg["src"]
    if not os.path.isfile(src):
        raise FileNotFoundError(f"找不到 Excel 文件: {src}")

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    if cfg["sheet"] not in wb.sheetnames:
        raise ValueError(f"工作簿里没有工作表「{cfg['sheet']}」，现有: {', '.join(wb.sheetnames)}")
    ws = wb[cfg["sheet"]]

    cs, ce = cfg["col_start"], cfg["col_end"]
    if ce < cs:
        raise ValueError("结束列不能小于起始列")
    ncols = ce - cs + 1
    split_col = cfg["split_col"]
    header_row = cfg["header_row"]

    sum_rels = []
    for sc in cfg["sum_cols"]:
        rel = sc - cs
        if 0 <= rel < ncols:
            sum_rels.append(rel)

    # ---- 表头
    headers = []
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        for j in range(cs, ce + 1):
            v = row[j - 1] if j - 1 < len(row) else None
            headers.append(fmt_cell(v) if fmt_cell(v) else f"列{index_to_col(j)}")
        break
    if not headers:
        headers = [f"列{index_to_col(j)}" for j in range(cs, ce + 1)]

    # ---- 分组
    groups = OrderedDict()
    skipped = 0
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        key = row[split_col - 1] if split_col - 1 < len(row) else None
        key_s = fmt_cell(key).strip()
        if not key_s:
            if cfg["empty_split"] == "skip":
                skipped += 1
                continue
            key_s = "未命名"
        vals = tuple((row[j - 1] if j - 1 < len(row) else None) for j in range(cs, ce + 1))
        groups.setdefault(key_s, []).append(vals)
    wb.close()

    out_dir = cfg["out_dir"]
    os.makedirs(out_dir, exist_ok=True)

    log(f"共读取到 {len(groups)} 个分组（拆分列 {index_to_col(split_col)}），开始生成...")
    if skipped:
        log(f"拆分列为空被跳过的行数: {skipped}")

    ok, fail, multi = 0, [], 0
    for i, (fname, rows) in enumerate(groups.items(), 1):
        try:
            _, pages = build_pdf(fname, headers, rows, sum_rels, out_dir,
                                 cfg["orientation"], cfg["title"],
                                 cfg["sum_label"], cfg["font_size"])
            ok += 1
            if pages > 1:
                multi += 1
        except Exception as exc:
            fail.append((fname, str(exc)))
            log(f"[失败] {fname}: {exc}")
        if i % 20 == 0 or i == len(groups):
            log(f"进度: {i}/{len(groups)}")

    log(f"完成！成功 {ok} 个，失败 {len(fail)} 个，其中多页文件 {multi} 个")
    log(f"输出目录: {out_dir}")
    return ok, fail, multi, len(groups)


# ---------------------------------------------------------------- 图形界面

def launch_gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    try:  # 高分屏清晰化
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass

    root = tk.Tk()
    root.title("Excel 按列拆分 PDF 工具")
    root.geometry("680x720")
    root.minsize(620, 660)

    main = ttk.Frame(root, padding=12)
    main.pack(fill="both", expand=True)

    def add_row(grid_row, label, widget, col2=None):
        ttk.Label(main, text=label).grid(row=grid_row, column=0, sticky="e", padx=(0, 6), pady=3)
        widget.grid(row=grid_row, column=1, sticky="ew", pady=3)
        if col2 is not None:
            col2.grid(row=grid_row, column=2, sticky="w", padx=(6, 0), pady=3)

    # --- 输入文件
    var_src = tk.StringVar()
    ent_src = ttk.Entry(main, textvariable=var_src)

    def pick_src():
        f = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm *.xls"), ("所有文件", "*.*")])
        if f:
            var_src.set(f)
            load_sheets()

    btn_src = ttk.Button(main, text="浏览…", command=pick_src)
    add_row(0, "Excel 文件：", ent_src, btn_src)

    # --- 工作表
    var_sheet = tk.StringVar()
    cmb_sheet = ttk.Combobox(main, textvariable=var_sheet, state="readonly")

    def load_sheets():
        try:
            wb = openpyxl.load_workbook(var_src.get(), read_only=True)
            cmb_sheet["values"] = list(wb.sheetnames)
            if wb.sheetnames:
                var_sheet.set(wb.sheetnames[0])
            wb.close()
        except Exception as exc:
            messagebox.showerror("读取失败", f"无法读取该 Excel：\n{exc}")

    btn_sheet = ttk.Button(main, text="刷新", command=load_sheets)
    add_row(1, "工作表：", cmb_sheet, btn_sheet)

    # --- 表头行 / 列范围 / 拆分列 / 合计列
    var_header_row = tk.IntVar(value=1)
    spn_header = ttk.Spinbox(main, from_=1, to=50, textvariable=var_header_row, width=8)
    add_row(2, "表头所在行：", spn_header)

    var_col_start = tk.StringVar(value="A")
    var_col_end = tk.StringVar(value="E")
    frame_cols = ttk.Frame(main)
    ttk.Entry(frame_cols, textvariable=var_col_start, width=6).pack(side="left")
    ttk.Label(frame_cols, text=" 列 到 ").pack(side="left")
    ttk.Entry(frame_cols, textvariable=var_col_end, width=6).pack(side="left")
    ttk.Label(frame_cols, text=" 列（填字母如 B、F）").pack(side="left", padx=(8, 0))
    add_row(3, "表格列范围：", frame_cols)

    var_split_col = tk.StringVar(value="F")
    ent_split = ttk.Entry(main, textvariable=var_split_col, width=8)
    add_row(4, "拆分依据列：", ent_split)
    ttk.Label(main, text="按该列的值分组，文件名也用该列的值",
              foreground="#666").grid(row=5, column=1, sticky="w")

    var_sum_cols = tk.StringVar()
    ent_sum = ttk.Entry(main, textvariable=var_sum_cols, width=20)
    add_row(6, "合计列：", ent_sum)
    ttk.Label(main, text="多个用逗号分隔（如 E,F），留空表示不需要合计行",
              foreground="#666").grid(row=7, column=1, sticky="w")

    # --- 拆分列为空的处理
    var_empty = tk.StringVar(value="unnamed")
    frame_empty = ttk.Frame(main)
    ttk.Radiobutton(frame_empty, text="跳过该行", variable=var_empty, value="skip").pack(side="left")
    ttk.Radiobutton(frame_empty, text="归入“未命名”文件", variable=var_empty, value="unnamed").pack(side="left", padx=(12, 0))
    add_row(8, "拆分列为空时：", frame_empty)

    # --- 纸张方向
    var_orient = tk.StringVar(value="landscape")
    frame_orient = ttk.Frame(main)
    ttk.Radiobutton(frame_orient, text="横向", variable=var_orient, value="landscape").pack(side="left")
    ttk.Radiobutton(frame_orient, text="纵向", variable=var_orient, value="portrait").pack(side="left", padx=(16, 0))
    add_row(9, "纸张方向：", frame_orient)

    # --- 标题 / 汇总标签 / 字号
    var_title = tk.StringVar(value="附件：")
    ent_title = ttk.Entry(main, textvariable=var_title)
    add_row(10, "PDF 标题：", ent_title)
    ttk.Label(main, text="显示在表格左上方，可留空", foreground="#666").grid(row=11, column=1, sticky="w")

    var_sum_label = tk.StringVar(value="汇总")
    ent_sumlabel = ttk.Entry(main, textvariable=var_sum_label, width=10)
    add_row(12, "汇总行标签：", ent_sumlabel)

    var_size = tk.DoubleVar(value=9.0)
    spn_size = ttk.Spinbox(main, from_=6, to=14, increment=0.5, textvariable=var_size, width=8)
    add_row(13, "字号(pt)：", spn_size)

    # --- 输出目录
    var_out = tk.StringVar()
    ent_out = ttk.Entry(main, textvariable=var_out)

    def pick_out():
        d = filedialog.askdirectory(title="选择输出文件夹")
        if d:
            var_out.set(d)

    btn_out = ttk.Button(main, text="浏览…", command=pick_out)
    add_row(14, "输出文件夹：", ent_out, btn_out)

    main.columnconfigure(1, weight=1)

    # --- 日志区
    ttk.Separator(main).grid(row=15, column=0, columnspan=3, sticky="ew", pady=(10, 6))
    txt_log = tk.Text(main, height=12, state="disabled", font=("Microsoft YaHei UI", 9))
    txt_log.grid(row=16, column=0, columnspan=3, sticky="nsew")
    main.rowconfigure(16, weight=1)

    q = queue.Queue()

    def log(msg):
        q.put(str(msg))

    def poll():
        try:
            while True:
                msg = q.get_nowait()
                if msg == "__DONE__":
                    btn_run.config(state="normal")
                    if q.get_nowait() == "askopen":
                        if messagebox.askyesno("完成", "生成结束，是否打开输出文件夹？"):
                            os.startfile(var_out.get())
                    continue
                txt_log.config(state="normal")
                txt_log.insert("end", msg + "\n")
                txt_log.see("end")
                txt_log.config(state="disabled")
        except queue.Empty:
            pass
        root.after(120, poll)

    def collect_cfg():
        cfg = {
            "src": var_src.get().strip(),
            "sheet": var_sheet.get(),
            "header_row": int(var_header_row.get()),
            "col_start": col_to_index(var_col_start.get()),
            "col_end": col_to_index(var_col_end.get()),
            "split_col": col_to_index(var_split_col.get()),
            "sum_cols": [col_to_index(x) for x in re.split(r"[，,、\s]+", var_sum_cols.get()) if x.strip()],
            "empty_split": var_empty.get(),
            "orientation": var_orient.get(),
            "title": var_title.get(),
            "sum_label": var_sum_label.get() or "汇总",
            "font_size": float(var_size.get()),
            "out_dir": var_out.get().strip(),
        }
        if not cfg["src"]:
            raise ValueError("请先选择 Excel 文件")
        if not cfg["sheet"]:
            raise ValueError("请选择工作表")
        if not cfg["out_dir"]:
            raise ValueError("请选择输出文件夹")
        return cfg

    def worker(cfg):
        try:
            run_split(cfg, log=log)
        except Exception as exc:
            log(f"[错误] {exc}")
        finally:
            log("__DONE__")
            q.put("askopen")

    def on_run():
        try:
            cfg = collect_cfg()
        except Exception as exc:
            messagebox.showerror("参数有误", str(exc))
            return
        txt_log.config(state="normal")
        txt_log.delete("1.0", "end")
        txt_log.config(state="disabled")
        btn_run.config(state="disabled")
        threading.Thread(target=worker, args=(cfg,), daemon=True).start()

    btn_run = ttk.Button(main, text="开始生成 PDF", command=on_run)
    btn_run.grid(row=17, column=0, columnspan=3, pady=(10, 0), ipadx=20, ipady=4)

    root.after(120, poll)
    root.mainloop()


if __name__ == "__main__":
    launch_gui()
